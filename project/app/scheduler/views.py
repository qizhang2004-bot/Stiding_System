# -*- coding: utf-8 -*-
import logging
import re
import threading
import time
from collections import defaultdict
from datetime import date, timedelta
from typing import Dict, List, Optional, Tuple

from django.contrib.auth import authenticate, login as auth_login, logout as auth_logout
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.db.models import Q
from django.shortcuts import get_object_or_404, redirect, render
from django.utils.http import url_has_allowed_host_and_scheme
from django.views.decorators.http import require_http_methods

from .models import Assignment, Group, Person, Role, Schedule, Team, UserProfile

# 排班算法模型（单一文件，见 scheduling.py 的模块说明）
from project.app.scheduler.scheduling import build_schedule, capacity_analysis, capacity_quick

# 周期起算日：25 号开始算下一个月（某月 M 的周期 = 上月25号 ~ 本月24号）
PERIOD_START_DAY = 25
FIXED_SHIFTS = ["早班", "中班", "晚班"]

# 操作审计日志（改班/导入/生成等敏感操作，输出到 scheduler.audit logger）
audit_log = logging.getLogger("scheduler.audit")

# 登录防爆破：同 IP+账号 5 次失败锁 5 分钟（进程内存级，够用即可）
_login_attempts: Dict[str, List[float]] = defaultdict(list)
LOGIN_MAX_ATTEMPTS = 5
LOGIN_LOCK_SECONDS = 300

# 生成排班互斥锁：同一班组同时只允许一个生成任务（避免并发重复求解/写库冲突）
_generate_locks: Dict[int, threading.Lock] = defaultdict(threading.Lock)


# ---------------------------------------------------------------------------
# 登录 / 团队权限
# ---------------------------------------------------------------------------
def login_view(request):
    if request.user.is_authenticated:
        return redirect("scheduler:index")
    error = ""
    if request.method == "POST":
        username = (request.POST.get("username") or "").strip()
        password = request.POST.get("password") or ""
        key = f"{request.META.get('REMOTE_ADDR', '?')}:{username}"
        now = time.monotonic()
        attempts = [t for t in _login_attempts.get(key, []) if now - t < LOGIN_LOCK_SECONDS]
        if len(attempts) >= LOGIN_MAX_ATTEMPTS:
            wait = max(1, int(LOGIN_LOCK_SECONDS - (now - attempts[0])))
            audit_log.warning("登录被限流 ip=%s user=%s", request.META.get("REMOTE_ADDR"), username)
            return render(request, "scheduler/login.html", {
                "error": f"尝试次数过多，请 {wait} 秒后再试。",
                "next": request.GET.get("next", ""),
            })
        user = authenticate(request, username=username, password=password)
        if user is not None:
            _login_attempts.pop(key, None)
            auth_login(request, user)
            audit_log.info("登录成功 user=%s", username)
            next_url = request.POST.get("next") or request.GET.get("next") or "/"
            if not url_has_allowed_host_and_scheme(
                next_url, allowed_hosts={request.get_host()},
                require_https=request.is_secure(),
            ):
                next_url = "/"
            return redirect(next_url)
        _login_attempts[key] = attempts + [now]
        audit_log.warning("登录失败 ip=%s user=%s", request.META.get("REMOTE_ADDR"), username)
        error = "账号或密码错误，请重试。"
    return render(request, "scheduler/login.html", {
        "error": error,
        "next": request.GET.get("next", ""),
    })


def logout_view(request):
    auth_logout(request)
    return redirect("scheduler:login")


def user_group(request):
    """当前登录用户所属队组（未绑定/超级管理员返回 None = 看所有队组）。"""
    if not request.user.is_authenticated:
        return None
    profile = getattr(request.user, "profile", None)
    return profile.group if profile else None


def user_role(request):
    """当前登录用户角色：super / team_admin / member。"""
    if not request.user.is_authenticated:
        return "member"
    profile = getattr(request.user, "profile", None)
    if request.user.is_superuser:
        return "super"
    return profile.role if profile else "member"


def can_edit_team(request, team) -> bool:
    """是否允许编辑该班组：超级管理员 或 该班组所属队组的队组管理员。"""
    role = user_role(request)
    if role == "super":
        return True
    ug = user_group(request)
    return role == "team_admin" and ug is not None and team is not None and team.group_id == ug.id


# ---------------------------------------------------------------------------
# 周期 / 日期工具
# ---------------------------------------------------------------------------
def period_range(year: int, month: int):
    """返回 (周期开始日期, 周期结束日期, 天数)。某月 M 的周期 = 上月25号 ~ 本月24号。"""
    if month == 1:
        start = date(year - 1, 12, PERIOD_START_DAY)
    else:
        start = date(year, month - 1, PERIOD_START_DAY)
    end = date(year, month, PERIOD_START_DAY - 1)
    days = (end - start).days + 1
    return start, end, days


def current_period(today: date = None):
    """今天属于哪个排班周期（按 25 号划分），返回 (year, month)。"""
    today = today or date.today()
    if today.day >= PERIOD_START_DAY:
        y, m = today.year, today.month + 1
        if m == 13:
            y, m = y + 1, 1
        return y, m
    return today.year, today.month


def month_shift(year: int, month: int, delta: int):
    """返回 (year, month) 偏移 delta 个月。"""
    m = month - 1 + delta
    y = year + m // 12
    return y, m % 12 + 1


def _post_int(value, default=None):
    """把请求里的 id 参数安全转成 int，失败返回 default。

    直接拿用户输入的字符串喂给 ORM 的 filter(id=...) 会抛
    ``ValueError: Field 'id' expected a number`` 导致 500，这里统一兜底。
    """
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def _parse_year_month(request):
    """从 GET 参数解析 (year, month)，非法或越界时回退到当前周期。

    越界的 y（负数 / >9999）或 m 传给 datetime.date 会抛 ValueError → 500，
    因此在入口处钳制。
    """
    try:
        y = int(request.GET.get("y", 0) or 0)
        m = int(request.GET.get("m", 0) or 0)
    except (TypeError, ValueError):
        y = m = 0
    if not (1 <= y <= 9999 and 1 <= m <= 12):
        return current_period()
    return y, m


def _assignment_matrix(schedule: Schedule):
    """从 Assignment 表一次查询聚合出排班明细矩阵（唯一数据源）。

    返回 (per_day, worker_counts, by_person)：
        per_day[day][shift] = [names]     day 为 int、shift 为 str
        worker_counts[name] = 上班天数     （只含有班的人）
        by_person[name][day] = shift      个人日历快速查询
    """
    shifts = schedule.shifts or FIXED_SHIFTS
    per_day = {d: {s: [] for s in shifts} for d in range(schedule.days)}
    worker_counts = {}
    by_person = {}
    for a in Assignment.objects.filter(schedule=schedule).select_related("person"):
        name = a.person.name
        per_day.setdefault(a.day, {}).setdefault(a.shift, []).append(name)
        worker_counts[name] = worker_counts.get(name, 0) + 1
        by_person.setdefault(name, {})[a.day] = a.shift
    return per_day, worker_counts, by_person


def _worked_auto_map(persons, schedule: Schedule = None) -> dict:
    """批量计算「已上班数」= 导入初始值 + 排班中日期已过的上班天数。

    一次查询该排班的所有明细，避免对每个人各查一次（N+1 查询）。
    返回 {person_id: 已上班数}。
    """
    result = {p.id: p.worked_so_far for p in persons}
    if not schedule or not schedule.start_date:
        return result
    today = date.today()
    days_by_person = {}
    for pid, day in Assignment.objects.filter(
        schedule=schedule, person__in=list(persons)
    ).values_list("person_id", "day"):
        days_by_person.setdefault(pid, []).append(day)
    for pid, days in days_by_person.items():
        result[pid] += sum(1 for d in days if schedule.start_date + timedelta(days=d) <= today)
    return result


def person_worked_auto(person: Person, schedule: Schedule = None) -> int:
    """已上班数（按日期自动计算）= 导入初始值 + 排班中「日期已过」的上班天数。"""
    count = person.worked_so_far
    if schedule and schedule.start_date:
        today = date.today()
        days = Assignment.objects.filter(
            schedule=schedule, person=person).values_list("day", flat=True)
        count += sum(1 for d in days if schedule.start_date + timedelta(days=d) <= today)
    return count


def _worked_auto_batch(pairs: List[Tuple[Person, Optional[Schedule]]]) -> Dict[int, int]:
    """批量计算「已上班数」（替代逐人 person_worked_auto，避免 N+1 查询）。

    pairs: [(人员, 其当月排班), ...]；返回 {person_id: 已上班数}。
    """
    result = {p.id: p.worked_so_far for p, _ in pairs}
    today = date.today()
    by_schedule: Dict[int, Tuple[Schedule, List[int]]] = {}
    for p, sch in pairs:
        if sch and sch.start_date:
            by_schedule.setdefault(sch.id, (sch, []))[1].append(p.id)
    for sch_id, (sch, pids) in by_schedule.items():
        for pid, day in Assignment.objects.filter(
            schedule_id=sch_id, person_id__in=pids
        ).values_list("person_id", "day"):
            if sch.start_date + timedelta(days=day) <= today:
                result[pid] += 1
    return result


def _xlsx_response(workbook, filename: str):
    """把 openpyxl Workbook 输出为 xlsx 下载响应（中文文件名用 RFC 5987 编码）。"""
    from io import BytesIO
    from urllib.parse import quote
    from django.http import HttpResponse
    buf = BytesIO()
    workbook.save(buf)
    buf.seek(0)
    resp = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(filename)}"
    return resp


def _export_board_xlsx(year: int, month: int, start, days: int, board: dict):
    """导出班次展示 Excel：第一行日期；早班/中班/晚班各一行，
    每个格子里按班组分类列出班组成员，班组之间空行隔开，文本自动换行。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook()
    ws = wb.active
    ws.title = "班次展示"
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    wrap = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=1, column=1, value="日期").font = bold
    for d in range(days):
        c = ws.cell(row=1, column=2 + d, value=f"{(start + timedelta(days=d)):%m-%d}")
        c.font = bold
        c.alignment = center
    for i, s in enumerate(FIXED_SHIFTS):
        row = 2 + i
        ws.cell(row=row, column=1, value=s).font = bold
        ws.cell(row=row, column=1).alignment = center
        for d in range(days):
            teams = board[d][s]
            # 班组与班组之间空行隔开
            text = "\n\n".join(
                f"{tname}：{'、'.join(names)}" for tname, names in teams.items() if names
            )
            cell = ws.cell(row=row, column=2 + d, value=text)
            cell.alignment = wrap
    ws.column_dimensions["A"].width = 8
    for col in range(2, days + 2):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 26
    return _xlsx_response(wb, f"班次展示_{year}年{month:02d}月.xlsx")


def _export_person_xlsx(person, year: int, month: int, start, days: int, assignments: dict):
    """导出个人日历 Excel：第一行日期，第二行该人当天是否上班（上班显示班次，休息显示「休息」）。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook()
    ws = wb.active
    ws.title = "个人日历"
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    wrap = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=1, column=1, value="日期").font = bold
    for d in range(days):
        c = ws.cell(row=1, column=2 + d, value=f"{(start + timedelta(days=d)):%m-%d}")
        c.font = bold
        c.alignment = center
    ws.cell(row=2, column=1, value=person.name).font = bold
    for d in range(days):
        cell = ws.cell(row=2, column=2 + d, value=assignments.get(d) or "休息")
        cell.alignment = wrap
    ws.column_dimensions["A"].width = 10
    for col in range(2, days + 2):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 8
    return _xlsx_response(wb, f"个人日历_{person.name}_{year}年{month:02d}月.xlsx")


# ---------------------------------------------------------------------------
# 首页
# ---------------------------------------------------------------------------
@login_required
def index(request):
    ug = user_group(request)
    role = user_role(request)
    groups = Group.objects.order_by("name")

    # 超管：?group= 点击切换队组展示（默认第一个）；队组管理员/队员固定本队组
    selected_group = None
    if role == "super":
        sel = request.GET.get("group", "")
        selected_group = groups.filter(id=sel).first() if sel.isdigit() else groups.first()
    else:
        selected_group = ug

    y, m = current_period()
    start, end, days = period_range(y, m)

    # 折线图数据：该队组所有班组每天各班次上班人数（当前周期）
    chart = {
        "dates": [(start + timedelta(days=d)).strftime("%m-%d") for d in range(days)],
        "early": [0] * days, "mid": [0] * days, "night": [0] * days, "total": [0] * days,
    }
    schedules = []
    team_count = person_count = 0
    cap_total = 0
    shortfall = []
    if selected_group:
        schedules = _latest_schedules_for_month(y, m, selected_group)
        for sch in schedules:
            per_day, _, _ = _assignment_matrix(sch)
            for d in range(min(days, sch.days)):
                for s in FIXED_SHIFTS:
                    n = len(per_day.get(d, {}).get(s, []))
                    if s == "早班":
                        chart["early"][d] += n
                    elif s == "中班":
                        chart["mid"][d] += n
                    elif s == "晚班":
                        chart["night"][d] += n
                    chart["total"][d] += n
        team_count = Team.objects.filter(group=selected_group).count()
        person_count = Person.objects.filter(team__group=selected_group, is_active=True).count()
        # 可达最低出勤总人数：各班组 min(启用人数, 每天人数×天数 // 最少班数) 求和
        for team in Team.objects.filter(group=selected_group):
            target = team.min_shift_target or 18
            active = Person.objects.filter(team=team, is_active=True).count()
            if team.daily_headcount > 0 and target > 0:
                cap_total += min(active, (team.daily_headcount * days) // target)
        # 未达最低出勤名单（除去豁免人员），基于最新排班实时计算
        for sch in schedules:
            exempt = set(sch.exempt_names or [])
            target_global = sch.min_shift_target or 0
            _, wc, _ = _assignment_matrix(sch)
            snap = {s2["name"]: s2 for s2 in (sch.worker_snapshot or [])}
            for nm, s2 in snap.items():
                if nm in exempt:
                    continue
                if target_global > 0:
                    tgt = target_global - s2.get("worked", 0)
                elif s2.get("required", 0) > 0:
                    tgt = s2["required"] - s2.get("worked", 0)
                else:
                    tgt = 0
                if tgt <= 0:
                    continue
                cnt = wc.get(nm, 0)
                if cnt < tgt:
                    shortfall.append({
                        "name": nm, "team": sch.team.name if sch.team else "—",
                        "target": tgt, "count": cnt, "gap": tgt - cnt,
                    })
        shortfall.sort(key=lambda x: x["target"] - x["count"], reverse=True)

    has_schedule = len(schedules) > 0

    return render(request, "scheduler/index.html", {
        "role": role,
        "user_group": ug,
        "groups": groups,
        "selected_group": selected_group,
        "year": y, "month": m,
        "start": start, "end": end, "days": days,
        "chart": chart,
        "has_schedule": has_schedule,
        "team_count": team_count,
        "person_count": person_count,
        "schedule_count": len(schedules),
        "cap_total": cap_total,
        "shortfall": shortfall,
        "recent": Schedule.objects.filter(team__group=selected_group).order_by("-created_at")[:5]
        if selected_group else [],
    })


# ---------------------------------------------------------------------------
# 班组管理（点击班组 → 显示存储约束 + 该班人员）
# ---------------------------------------------------------------------------
def _parse_import_text(text: str, group_names=None):
    """解析导入文本，返回 [(队组, 班组, 姓名, [岗位...], 已上班数, 应上班数, 默认班次), ...]。

    每行两种格式：
      ① 班组-姓名-岗位1,岗位2-已上班数-应上班数-默认班次          （队组账号用）
      ② 队组-班组-姓名-岗位1,岗位2-已上班数-应上班数-默认班次      （超管用，最前方加队组）
      - 默认班次可写 早班/中班/晚班，写在最后、可省略（省略则保持默认「早班」）
    """
    group_names = group_names or set()
    items = []
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        if ":" in line or "：" in line:
            head = re.split(r"[:：]", line, maxsplit=1)
            role = head[0].strip()
            names = re.split(r"[,，、\s]+", head[1].strip()) if len(head) > 1 else []
            for nm in names:
                if nm:
                    items.append(("", nm, [role], 0, 0, ""))
            continue
        if "-" in line:
            parts = [p.strip() for p in line.split("-") if p.strip()]
        else:
            parts = [p.strip() for p in re.split(r"[,，、\s]+", line) if p.strip()]
        if not parts:
            continue
        # 从末尾抠默认班次
        default_shift = ""
        if parts[-1] in ("早班", "中班", "晚班"):
            default_shift = parts.pop()
        # 从末尾抠数字（已上班数、应上班数；允许 "0,18" 或 "0 18"）
        counts = []  # 最终为 [已上班数, 应上班数]
        while parts and len(counts) < 2:
            last_tokens = [t for t in re.split(r"[,，、\s]+", parts[-1]) if t.strip()]
            if len(last_tokens) >= 2 and all(re.fullmatch(r"\d+", t) for t in last_tokens[:2]):
                counts = [int(last_tokens[0]), int(last_tokens[1])]
                parts.pop()
                break
            if len(last_tokens) == 1 and re.fullmatch(r"\d+", last_tokens[0]):
                counts.insert(0, int(last_tokens[0]))
                parts.pop()
                continue
            break
        worked, required = (counts[0], counts[1]) if len(counts) == 2 else \
                           (0, counts[0]) if len(counts) == 1 else (0, 0)
        if "-" in line:
            first = parts.pop(0)
        else:
            first = ""
        # 第一个字段若是已知队组名，则它是「队组」，第二个字段才是「班组」
        group = ""
        if first in group_names:
            group = first
            team = parts.pop(0) if parts else ""
        else:
            team = first
        if not parts:
            # 只写了「姓名-数字」没有班组（如 王五-0-18）：第一段当姓名，
            # 班组留空 → 导入时默认归入当前选中的班组
            if team and team not in group_names:
                items.append((group, "", team, [], worked, required, default_shift))
            continue
        name = parts.pop(0)
        roles = []
        for p in parts:
            roles.extend([r.strip() for r in re.split(r"[,，、]+", p) if r.strip()])
        items.append((group, team, name, roles, worked, required, default_shift))
    return items


def _rows_from_excel(raw_bytes):
    """从 Excel（.xlsx/.xls）二进制解析出人员文本行（每行单元格用 - 连接）。"""
    from io import BytesIO
    import openpyxl
    wb = openpyxl.load_workbook(BytesIO(raw_bytes), read_only=True, data_only=True)
    ws = wb.active
    lines = []
    for row in ws.iter_rows(values_only=True):
        cells = [str(c).strip() if c is not None else "" for c in row]
        cells = [c for c in cells if c]
        if cells:
            lines.append("-".join(cells))
    return lines


def _rows_from_csv(text):
    """从 CSV 文本解析出人员文本行（每行单元格用 - 连接）。"""
    import csv
    import io
    lines = []
    for row in csv.reader(io.StringIO(text)):
        cells = [c.strip() for c in row if c.strip()]
        if cells:
            lines.append("-".join(cells))
    return lines


def _claim_account(g2: Group, role: str, username: str, password: str):
    """新增队组时认领账号：账号不存在则创建；存在但无任何队组绑定（孤儿账号）则直接复用；
    已绑定其他队组则拒绝并说明绑定关系。返回错误信息或 None。"""
    existing = User.objects.filter(username=username).first()
    if existing:
        prof = getattr(existing, "profile", None)
        if prof is not None:
            gname = prof.group.name if prof.group else "未绑定队组"
            return f"账号「{username}」已存在（绑定队组「{gname}」），请换一个账号名。"
        # 孤儿账号（无任何绑定）：复用，重设密码并绑到新队组
        existing.set_password(password or "111111")
        existing.save()
        UserProfile.objects.create(user=existing, group=g2, role=role)
        return None
    u = User.objects.create_user(username, password=password or "111111")
    UserProfile.objects.create(user=u, group=g2, role=role)
    return None


def _upsert_group_account(g2: Group, role: str, username: str, password: str):
    """编辑队组时更新/创建绑定账号：用户名可改；密码留空=保持原密码。返回错误信息或 None。"""
    profile = UserProfile.objects.filter(group=g2, role=role).select_related("user").first()
    if profile and profile.user:
        u = profile.user
        if User.objects.filter(username=username).exclude(id=u.id).exists():
            other = User.objects.filter(username=username).exclude(id=u.id).first()
            oprof = getattr(other, "profile", None)
            oname = oprof.group.name if oprof and oprof.group else "未绑定队组"
            return f"账号「{username}」已存在（绑定队组「{oname}」），请换一个账号名。"
        u.username = username
        if password:
            u.set_password(password)
        u.save()
        return None
    if User.objects.filter(username=username).exists():
        return f"账号「{username}」已存在，请换一个账号名。"
    u = User.objects.create_user(username, password=password or "111111")
    UserProfile.objects.create(user=u, group=g2, role=role)
    return None


def _int_or(value, default: int) -> int:
    """把表单值安全转成 int，失败返回 default。"""
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return default


def _parse_person_overrides(request, team) -> dict:
    """解析「专人专项约束」表单，返回 {姓名: {"rest_min":int,"rest_max":int,"work_days":int}}。

    - 只保留本班组真实存在的启用人员（防止越权/脏数据）
    - 同一人重复填写时以最后一行为准
    - 姓名留空的行直接忽略（前端「点击添加一行」会新增空行）
    - 连休最小值强制 ≥ 2（禁止单休），上限不小于下限
    - 最少上班天数 0 表示该项沿用班组默认
    """
    names = request.POST.getlist("ov_person")
    rmins = request.POST.getlist("ov_rest_min")
    rmaxs = request.POST.getlist("ov_rest_max")
    works = request.POST.getlist("ov_work_days")
    if not names:
        return {}
    valid = set(Person.objects.filter(
        team=team, is_active=True).values_list("name", flat=True))
    out = {}
    for i, raw_name in enumerate(names):
        nm = (raw_name or "").strip()
        if not nm or nm not in valid:
            continue
        rmin = max(2, _int_or(rmins[i] if i < len(rmins) else None, 2))
        rmax = max(rmin, _int_or(rmaxs[i] if i < len(rmaxs) else None, rmin))
        work = max(0, _int_or(works[i] if i < len(works) else None, 0))
        out[nm] = {"rest_min": rmin, "rest_max": rmax, "work_days": work}
    return out


def _save_team_constraints(request, team) -> Tuple[str, str]:
    """把「排班约束」表单写入班组。返回 (warn 片段, 错误信息)。

    含字段级审核：数字非法、岗位条件 op 白名单、连休最小值钳到 2、
    豁免名单只保留本班组真实人员、专人专项只保留本班组启用人员。
    """
    try:
        daily_headcount = max(0, int(request.POST.get("daily_headcount") or 0))
    except ValueError:
        return "", "「该班应上人数（每天）」必须是整数。"
    role_names = request.POST.getlist("role_names")
    role_ops = request.POST.getlist("role_ops")
    role_counts = request.POST.getlist("role_counts")
    role_reqs = {}
    for rn, op, cnt in zip(role_names, role_ops, role_counts):
        rn = (rn or "").strip()
        if not rn or len(rn) > 50:
            continue
        op = op if op in (">=", "<=", "==") else ">="
        try:
            count = max(0, int(cnt or 0))
        except (TypeError, ValueError):
            count = 0
        role_reqs[rn] = {"op": op, "count": count}
    try:
        # 需求：休息至少 2 天（禁止单休），强制下限为 2
        rb_min = max(2, int(request.POST.get("rest_min") or 2))
        rb_max = max(rb_min, int(request.POST.get("rest_max") or 4))
    except ValueError:
        rb_min, rb_max = 2, 4
    try:
        min_shift_target = max(0, int(request.POST.get("min_shift_target") or 0))
    except ValueError:
        return "", "「每人应上最少班数」必须是整数。"
    # 豁免名单：只保留本班组真实存在的启用人员（按勾选顺序去重）
    exempt_selected = request.POST.getlist("exempt")
    if exempt_selected:
        valid = set(Person.objects.filter(
            team=team, name__in=exempt_selected, is_active=True
        ).values_list("name", flat=True))
        exempt_names = [n for n in exempt_selected if n in valid]
    else:
        exempt_names = []
    # 专人专项约束
    person_overrides = _parse_person_overrides(request, team)

    team.daily_headcount = daily_headcount
    team.role_reqs = role_reqs
    team.rest_block = {"min": rb_min, "max": rb_max}
    team.min_shift_target = min_shift_target
    team.exempt_names = exempt_names
    team.person_overrides = person_overrides
    team.save()
    audit_log.info(
        "保存约束 team=%s daily=%s roles=%s rest=%s target=%s exempt=%s overrides=%s user=%s",
        team.name, daily_headcount, role_reqs, team.rest_block,
        min_shift_target, exempt_names, person_overrides, request.user.username,
    )
    active_count = Person.objects.filter(team=team, is_active=True).count()
    warn = "&warn=daily" if daily_headcount > active_count else ""
    return warn, ""


@require_http_methods(["GET", "POST"])
@login_required
def team_manage(request):
    ug = user_group(request)
    groups = Group.objects.order_by("name")

    # 选择队组：队组账号固定为自己的队组；超级管理员通过 ?group=<id> 选择
    sel_group = request.GET.get("group", "")
    if ug:
        selected_group = ug
    else:
        selected_group = groups.filter(id=sel_group).first() if sel_group.isdigit() else None

    teams = Team.objects.select_related("group").order_by("name")
    if selected_group:
        teams = teams.filter(group=selected_group)
    else:
        teams = Team.objects.none()

    # 通过 ?team=<id 或 名称> 选择班组
    sel = request.GET.get("team", "")
    default_team = None
    if sel:
        default_team = teams.filter(id=sel).first() if sel.isdigit() else teams.filter(name=sel).first()
    default_team = default_team or teams.first()

    gq = f"&group={selected_group.id}" if selected_group else ""  # 跳转时保留队组选择

    message = ""
    error = ""
    if request.GET.get("deleted"):
        message = f"已删除班组「{request.GET['deleted']}」。"
    if request.GET.get("renamed") and default_team:
        message = f"班组已重命名为「{request.GET['renamed']}」。"
    if request.GET.get("added_group"):
        message = "已新增队组，队组管理员与队员账号已创建。"
    if request.GET.get("edited_group"):
        message = "队组信息与账号已更新。"
    if request.GET.get("deleted_group"):
        message = f"已删除队组「{request.GET['deleted_group']}」（其下班组与绑定账号一并删除）。"
    if request.GET.get("saved") and default_team:
        message = f"已保存「{default_team.name}」的排班约束。"
    if request.GET.get("error") == "daily":
        error = "请先填写「该班应上人数（每天）」（大于 0）再生成排班。"
    if request.GET.get("warn") == "daily":
        error = "⚠️ 每天应上人数大于该班组启用人数，无法排班（每人每天最多上 1 班）。请降低每天人数或增加启用人员。"
    if request.GET.get("busy"):
        error = "该班组正在生成排班，请稍候再试。"

    if request.method == "POST":
        action = request.POST.get("action", "")
        if user_role(request) == "member":
            error = "队员账号为只读，只能查看，不能修改排班数据。"
        elif action == "add_group":
            # 超级管理员新增队组：队组 + 队组管理员账号 + 队员只读账号（密码不做强度校验）
            if user_role(request) != "super":
                error = "只有超级管理员可以新增队组。"
            else:
                gname = (request.POST.get("group_name") or "").strip()
                auser = (request.POST.get("admin_username") or "").strip()
                apwd = request.POST.get("admin_password") or ""
                muser = (request.POST.get("member_username") or "").strip()
                mpwd = request.POST.get("member_password") or ""
                if not gname or len(gname) > 50:
                    error = "请填写队组名称（不超过 50 字）。"
                elif not auser or not apwd:
                    error = "请填写队组管理员账号和密码。"
                elif not muser or not mpwd:
                    error = "请填写队员查看账号和密码。"
                elif auser == muser:
                    error = "队组管理员账号与队员账号不能相同。"
                elif Group.objects.filter(name=gname).exists():
                    error = f"队组名称「{gname}」已存在，请换一个。"
                else:
                    g2 = Group.objects.create(name=gname)
                    err = _claim_account(g2, "team_admin", auser, apwd)
                    if err:
                        g2.delete()
                        error = err
                    else:
                        err2 = _claim_account(g2, "member", muser, mpwd)
                        if err2:
                            g2.delete()
                            error = err2
                        else:
                            audit_log.info("新增队组 %s admin=%s member=%s user=%s",
                                           gname, auser, muser, request.user.username)
                            return redirect(f"{request.path}?group={g2.id}&added_group=1")

        elif action == "edit_group":
            # 超级管理员编辑队组：名称/缩写/管理员与队员账号（密码留空=不改密码）
            gid = _post_int(request.POST.get("group_id"))
            g2 = Group.objects.filter(id=gid).first() if gid else None
            if user_role(request) != "super":
                error = "只有超级管理员可以编辑队组。"
            elif not g2:
                error = "队组不存在或已被删除。"
            else:
                gname = (request.POST.get("group_name") or "").strip()
                auser = (request.POST.get("admin_username") or "").strip()
                apwd = request.POST.get("admin_password") or ""
                muser = (request.POST.get("member_username") or "").strip()
                mpwd = request.POST.get("member_password") or ""
                if not gname or len(gname) > 50:
                    error = "请填写队组名称（不超过 50 字）。"
                elif not auser or not muser:
                    error = "请填写队组管理员账号与队员账号。"
                elif auser == muser:
                    error = "队组管理员账号与队员账号不能相同。"
                elif Group.objects.filter(name=gname).exclude(id=g2.id).exists():
                    error = f"队组名称「{gname}」已存在，请换一个。"
                else:
                    err = _upsert_group_account(g2, "team_admin", auser, apwd)
                    if err:
                        error = err
                    err2 = _upsert_group_account(g2, "member", muser, mpwd) if not error else None
                    if err2:
                        error = err2
                if not error:
                    g2.name = gname
                    g2.save()
                    audit_log.info("编辑队组 %s admin=%s member=%s user=%s",
                                   gname, auser, muser, request.user.username)
                    return redirect(f"{request.path}?group={g2.id}&edited_group=1")

        elif action == "delete_group":
            # 超级管理员删除队组：需输入「删除」二次确认；连带删除其下班组与绑定账号
            gid = _post_int(request.POST.get("group_id"))
            confirm = (request.POST.get("confirm_text") or "").strip()
            g2 = Group.objects.filter(id=gid).first() if gid else None
            if user_role(request) != "super":
                error = "只有超级管理员可以删除队组。"
            elif not g2:
                error = "队组不存在或已被删除。"
            elif confirm != "删除":
                error = "确认失败：请输入「删除」两个字才能删除队组。"
            else:
                gname = g2.name
                User.objects.filter(profile__group=g2, is_superuser=False).delete()
                Team.objects.filter(group=g2).delete()
                g2.delete()
                audit_log.info("删除队组 %s user=%s", gname, request.user.username)
                from urllib.parse import quote
                return redirect(f"{request.path}?deleted_group={quote(gname)}")

        elif action in ("save_constraints", "generate"):
            tid = _post_int(request.POST.get("team_id"))
            team = teams.filter(id=tid).first() if tid else None
            if not team:
                error = "请先选择一个班组。"
            else:
                warn, err = _save_team_constraints(request, team)
                if err:
                    error = err
                elif action == "save_constraints":
                    return redirect(f"{request.path}?team={team.id}&saved=1{warn}{gq}")
                else:
                    # 生成排班：同一班组只允许一个生成任务，防止并发重复求解
                    lock = _generate_locks[team.id]
                    if not lock.acquire(blocking=False):
                        return redirect(f"{request.path}?team={team.id}&busy=1{gq}")
                    try:
                        return _run_generate(request, team)
                    finally:
                        lock.release()

        elif action == "batch_shift":
            # 批量修改默认班次（用于中班/夜班倒班时统一切换）
            pids = [pid for pid in (_post_int(v) for v in request.POST.getlist("person_ids")) if pid is not None]
            new_shift = request.POST.get("batch_shift", "")
            if new_shift not in FIXED_SHIFTS:
                error = "无效的班次。"
            elif not pids:
                error = "请先勾选要修改的人员。"
            else:
                qs = Person.objects.filter(id__in=pids)
                if ug:
                    qs = qs.filter(team__group=ug)
                n = qs.update(default_shift=new_shift)
                audit_log.info("批量改默认班次 user=%s count=%s shift=%s",
                               request.user.username, n, new_shift)
                message = f"已把 {n} 名人员的默认班次改为「{new_shift}」。"
                return redirect(request.get_full_path())

        elif action == "import":
            text = request.POST.get("import_text", "")
            uploaded = request.FILES.get("import_file")
            parse_error = ""
            if uploaded:
                fname = (uploaded.name or "").lower()
                raw = uploaded.read()
                try:
                    if fname.endswith((".xlsx", ".xls", ".xlsm")):
                        # Excel：按单元格读，每行拼成「班组-姓名-岗位-已上-应上-班次」
                        text = "\n".join(_rows_from_excel(raw))
                    elif fname.endswith(".csv"):
                        # CSV：按逗号分列，每行拼成同样格式
                        csv_text = raw.decode("utf-8-sig", errors="ignore")
                        text = "\n".join(_rows_from_csv(csv_text))
                    else:
                        # 纯文本（txt / 粘贴）
                        text = raw.decode("utf-8-sig", errors="ignore")
                except Exception as e:  # noqa: BLE001
                    parse_error = f"文件解析失败：{e}"
                    text = ""
            group_names = set(Group.objects.values_list("name", flat=True))
            parsed = _parse_import_text(text, group_names)
            if parse_error:
                error = parse_error
            elif not parsed:
                error = "没有解析到任何人员，请检查导入格式。"
            else:
                created_p = 0
                updated_p = 0
                row_errors = []
                for grp, team, nm, roles, worked, required, default_shift in parsed:
                    # ---- 字段级审核 ----
                    nm = (nm or "").strip()
                    if not nm or len(nm) > 50:
                        row_errors.append(f"{nm or '(空)'}: 姓名无效（非空且不超过 50 字）")
                        continue
                    roles = list(dict.fromkeys(
                        r for r in roles if r and (r or "").strip() and len(r) <= 50
                    ))
                    # 没有专门岗位的按「普通」处理（提示用户填写普通，未填时自动兜底）
                    if not roles:
                        roles = ["普通"]
                    worked = max(0, min(999, int(worked or 0)))
                    required = max(0, min(999, int(required or 0)))
                    if default_shift not in FIXED_SHIFTS:
                        default_shift = ""
                    # 归属队组：文本里写了队组 > 账号自己的队组 > URL 指定队组 > 空
                    if grp:
                        g = Group.objects.filter(name=grp).first()
                    else:
                        g = ug or selected_group
                    try:
                        person, is_new = Person.objects.get_or_create(name=nm)
                        # 跨队组保护：队组账号不能把其它队组的人员划到自己名下
                        if not is_new and ug and person.team and person.team.group_id \
                                and person.team.group_id != ug.id:
                            row_errors.append(f"{nm}: 已属于其它队组「{person.team.group.name}」，跳过")
                            continue
                        if is_new:
                            created_p += 1
                        else:
                            updated_p += 1
                        # 归属班组：文本写了班组名 → 与已有班组比较，同名放入、不同名创建新班组；
                        # 没写班组名 → 默认导入到当前选中的班组
                        if team:
                            tname = team
                            if g:
                                person.team, _ = Team.objects.get_or_create(group=g, name=tname)
                            else:
                                person.team = Team.objects.filter(name=tname).first()
                        else:
                            person.team = default_team
                        person.worked_so_far = worked
                        person.required_shifts = required
                        # 默认班次：显式写了就用；没写保持默认「早班」
                        if default_shift:
                            person.default_shift = default_shift
                        person.save()
                        for rn in roles:
                            role, _ = Role.objects.get_or_create(name=rn)
                            person.roles.add(role)
                    except Exception as e:  # noqa: BLE001 —— 单行出错不中断整体导入
                        row_errors.append(f"{nm}: {e}")
                audit_log.info("导入人员 user=%s 新增=%s 更新=%s 失败=%s",
                               request.user.username, created_p, updated_p, len(row_errors))
                message = (f"导入完成：新增 {created_p} 人、更新 {updated_p} 人"
                           f"{'（已归入本队组「' + ug.name + '」的班组）' if ug else ''}"
                           f"，共处理 {len(parsed)} 条记录。")
                if row_errors:
                    error = "部分人员导入失败：" + "；".join(row_errors[:5])

        elif action == "delete":
            pid = _post_int(request.POST.get("person_id"))
            if pid:
                qs = Person.objects.filter(id=pid)
                if ug:
                    qs = qs.filter(team__group=ug)
                deleted = list(qs.values_list("name", flat=True))
                qs.delete()
                audit_log.info("删除人员 %s user=%s", deleted, request.user.username)
                message = "已删除该人员。"

        elif action == "rename_team":
            tid = _post_int(request.POST.get("team_id"))
            new_name = (request.POST.get("team_name") or "").strip()
            team = teams.filter(id=tid).first() if tid else None
            if not team:
                error = "班组不存在或已被删除。"
            elif not can_edit_team(request, team):
                error = "只能重命名本队组的班组。"
            elif not new_name:
                error = "请输入新班组名称。"
            elif len(new_name) > 50:
                error = "班组名称过长（最多 50 字）。"
            elif Team.objects.filter(group=team.group, name=new_name).exclude(id=team.id).exists():
                error = f"该队组已有班组「{new_name}」，请换一个名称。"
            else:
                old_name = team.name
                team.name = new_name
                team.save()
                audit_log.info("重命名班组 %s -> %s user=%s",
                               old_name, new_name, request.user.username)
                from urllib.parse import quote
                return redirect(f"{request.path}?team={team.id}&renamed={quote(new_name)}{gq}")

        elif action == "add_team":
            tn = (request.POST.get("team_name") or "").strip()
            # 队组管理员新增的班组自动归入自己的队组；超管需要先选队组
            target_group = ug or selected_group
            if not target_group:
                error = "请先选择一个队组，再新增班组。"
            elif not tn:
                error = "请输入班组名称。"
            elif len(tn) > 50:
                error = "班组名称过长（最多 50 字）。"
            else:
                Team.objects.get_or_create(group=target_group, name=tn)
                audit_log.info("添加班组 %s group=%s user=%s",
                               tn, target_group.name, request.user.username)
                message = f"已添加班组「{tn}」。"
                return redirect(f"{request.path}?team={tn}{gq}")

        elif action == "delete_team":
            tid = _post_int(request.POST.get("team_id"))
            confirm = (request.POST.get("confirm_text") or "").strip()
            team = Team.objects.filter(id=tid).first() if tid else None
            if not team:
                error = "班组不存在或已被删除。"
            elif ug and team.group_id != ug.id:
                error = "只能删除本队组的班组。"
            elif confirm != "删除":
                error = "确认失败：请输入「删除」两个字才能删除班组。"
            else:
                tname = team.name
                # 删除班组：其人员变为未分组，排班记录保留但失去班组归属
                team.delete()
                audit_log.info("删除班组 %s user=%s", tname, request.user.username)
                from urllib.parse import quote
                return redirect(f"{request.path}?deleted={quote(tname)}{gq}")

    # 该班人员（含按日期自动计算的已上班数）
    persons = Person.objects.filter(team=default_team).select_related("team").prefetch_related("roles").order_by("name") \
        if default_team else Person.objects.none()
    # 批量计算已上班数：取该班组最新的一份排班（order_by -created_at 后取第一条）
    latest_schedule = Schedule.objects.filter(team=default_team).order_by("-created_at").first() \
        if default_team else None
    persons_list = list(persons)
    worked_map = _worked_auto_map(persons_list, latest_schedule) if latest_schedule else {}
    rows = [(p, worked_map.get(p.id, p.worked_so_far)) for p in persons_list]

    # 固定岗位下拉选项：该班所有人的岗位 去重后的唯一集合（元组），并合并已配置的岗位
    if default_team:
        person_roles = set()
        for p in persons_list:
            person_roles.update(p.roles.values_list("name", flat=True))
        person_roles.update((default_team.role_reqs or {}).keys())
        team_role_options = tuple(sorted(person_roles))
    else:
        team_role_options = ()

    # 连休范围默认显示：取该班组已保存的配置；未保存过时与表单默认值保持一致。
    y0, m0 = current_period()
    _, _, period_days = period_range(y0, m0)
    if default_team and (default_team.min_shift_target or 0) > 0:
        cap_target = default_team.min_shift_target
    else:
        reqs = [p.required_shifts for p in persons_list if p.required_shifts > 0]
        cap_target = (sum(reqs) // len(reqs)) if reqs else 18
    rb_defaults = (default_team.rest_block if default_team else None) or {}
    rest_defaults = {
        "min": max(2, _int_or(rb_defaults.get("min"), 2)),
        "max": max(2, _int_or(rb_defaults.get("max"), 5)),
    }
    default_rest_max = rest_defaults["max"]
    # 专人专项约束：只展示仍然存在的启用人员（人员被删/禁用后自动不再展示）
    active_names = {p.name for p in persons_list if p.is_active}
    override_rows = [
        {"name": nm,
         "rest_min": max(2, _int_or(spec.get("rest_min"), 2)),
         "rest_max": max(2, _int_or(spec.get("rest_max"), 5)),
         "work_days": max(0, _int_or(spec.get("work_days"), 0))}
        for nm, spec in ((default_team.person_overrides or {}) if default_team else {}).items()
        if nm in active_names and isinstance(spec, dict)
    ]

    # 容量预估（提示最多能有多少人排满，需要豁免几人；只统计启用人员）
    # 关键：必须先给「专人专项」的人留够他们要求的班次，剩下的才能分给普通人，
    # 否则会忽略专项占用、把豁免人数算少（例如 25 人 3 个专项 19 班时，
    # 忽略专项会算出「只差 1 人」，实际是差 2 人）。
    capacity = None
    special_days = [r["work_days"] for r in override_rows if r["work_days"] > 0]
    capacity_group = None
    if default_team:
        people_count = Person.objects.filter(team=default_team, is_active=True).count()
        exempt_count = len([n for n in (default_team.exempt_names or []) if
                            Person.objects.filter(team=default_team, name=n, is_active=True).exists()])
        # 已豁免的人不参与「最少班数」约束，也不算进专项占用
        exempt_active = {n for n in (default_team.exempt_names or [])}
        special_days = [r["work_days"] for r in override_rows
                        if r["work_days"] > 0 and r["name"] not in exempt_active]
        capacity = capacity_analysis(
            people_count, default_team.daily_headcount or 0, period_days,
            default_rest_max, cap_target, exempt_count,
            special_days=special_days,
        )
        capacity_group = {
            "totalPeople": people_count,
            "daily": default_team.daily_headcount or 0,
            "special": [{"name": r["name"], "days": r["work_days"]} for r in override_rows],
            "exempt": sorted(exempt_active),
            "target": cap_target,
        }

    # 每个岗位的持有人数（用于实时"岗位条件可行性"检查；只统计启用人员）
    role_holder_counts = {}
    if default_team:
        for p in persons_list:
            if not p.is_active:
                continue
            for rn in p.roles.values_list("name", flat=True):
                role_holder_counts[rn] = role_holder_counts.get(rn, 0) + 1

    # 启用人数（参与排班的人数，前端「约束实时自检」用它算需豁免人数）
    active_person_count = Person.objects.filter(team=default_team, is_active=True).count() if default_team else 0

    return render(request, "scheduler/team_manage.html", {
        "teams": teams,
        "team": default_team,
        "user_group": ug,
        "groups": groups,
        "selected_group": selected_group,
        "shifts": FIXED_SHIFTS,
        "rows": rows,
        "team_persons": [p for p, _ in rows],
        "active_person_count": active_person_count,
        "team_role_options": team_role_options,
        "role_holder_counts": role_holder_counts,
        "default_rest_max": default_rest_max,
        "rest_defaults": rest_defaults,
        "override_rows": override_rows,
        "active_person_names": sorted(active_names),
        "capacity_group": capacity_group,
        "exempt_names_json": sorted(exempt_active) if default_team else [],
        "capacity": capacity,
        "period_days": period_days,
        "is_team_user": ug is not None,
        "is_super": user_role(request) == "super",
        "can_manage": user_role(request) in ("super", "team_admin"),
        "can_edit": can_edit_team(request, default_team),
        "roles": Role.objects.order_by("name"),
        "message": message,
        "error": error,
    })


def _run_generate(request, team: Team):
    """用班组存储的约束调用引擎生成排班，返回重定向到结果页。"""
    if not team.daily_headcount or team.daily_headcount <= 0:
        return redirect(f"/teams/?group={team.group_id or ''}&team={team.id}&error=daily")
    audit_log.info("开始生成排班 team=%s user=%s", team.name, request.user.username)
    y, m = current_period()
    start, end, days = period_range(y, m)
    persons = Person.objects.filter(team=team, is_active=True).prefetch_related("roles").order_by("name")
    if not persons:
        return redirect(f"/teams/?group={team.group_id or ''}&team={team.id}")
    worker_snapshot = [
        {"name": p.name, "roles": list(p.roles.values_list("name", flat=True)),
         "worked": p.worked_so_far, "required": p.required_shifts,
         "team": team.name, "default_shift": p.default_shift}
        for p in persons
    ]
    default_shift_map = {}
    exempt_set = set(team.exempt_names or [])
    target_global = team.min_shift_target or 0
    for p in persons:
        default_shift_map[p.name] = p.default_shift


    # 容量预估的 target：全局「至少应上班数」优先；否则用每人「应上班数」的平均
    if target_global > 0:
        cap_target = target_global
    else:
        reqs = [p.required_shifts for p in persons if p.required_shifts > 0]
        cap_target = (sum(reqs) // len(reqs)) if reqs else 18
    # 连休范围：取班组存储的默认值（在班组管理界面「1.3 连休范围」里配置）；
    # 未设置时回退到与表单默认一致的 2~5 天。
    # 连续上班：至少 2 天（禁止只上一天班就休息）；
    # 上限 = (周期天数 // 3) − 最短连续休息天数。
    eff_target = cap_target if cap_target > 0 else 18
    rb_stored = team.rest_block or {}
    rest_min = max(2, _int_or(rb_stored.get("min"), 2))
    rest_max = max(rest_min, _int_or(rb_stored.get("max"), 5))
    work_run_min = 2
    work_run_max = max(2, days // 3 - rest_min)

    # 容量预估：最多能有多少人排满目标（豁免人员不参与休息计算，不占最少班数）
    non_exempt_count = len([p for p in persons if p.name not in exempt_set])
    cap = capacity_analysis(
        len(persons), team.daily_headcount or 0, days,
        rest_max, cap_target,
        exempt_count=len(persons) - non_exempt_count,
    )

    # 专人专项约束：把班组存储的 {姓名: {rest_min, rest_max, work_days}}
    # 转成引擎的 worker_rules（该人的要求完全取代班组默认值）。
    # 豁免人员不参与休息规则，其专项约束不生效，跳过。
    person_overrides = {
        nm: spec for nm, spec in (team.person_overrides or {}).items()
        if nm not in exempt_set and isinstance(spec, dict)
    }
    worker_rules = {}
    for nm, spec in person_overrides.items():
        r_min = max(2, _int_or(spec.get("rest_min"), rest_min))
        r_max = max(r_min, _int_or(spec.get("rest_max"), rest_max))
        rule = {"rest_block": {"min": r_min, "max": r_max}}
        work_days = max(0, _int_or(spec.get("work_days"), 0))
        if work_days > 0:
            rule["work_days"] = work_days
        worker_rules[nm] = rule

    base_config = {
        "workers": worker_snapshot,
        "shifts": FIXED_SHIFTS,
        "days": days,
        "daily_total": team.daily_headcount or 0,
        "role_req": team.role_reqs or {},
        "min_shift_target": team.min_shift_target or None,
        "worker_default_shift": default_shift_map,
        "exempt_workers": team.exempt_names or [],
        "rest_block": {"min": rest_min, "max": rest_max},
        "work_block": {"min": work_run_min, "max": work_run_max},
        "worker_rules": worker_rules,
    }

    def _build_req():
        req = {}
        has_exempt = bool(exempt_set)
        for p in persons:
            # 规则3：全局「至少应上班数」优先；未填时才用每人「应上班数」
            if target_global > 0:
                tgt = target_global - p.worked_so_far
            elif p.required_shifts > 0:
                tgt = p.required_shifts - p.worked_so_far
            else:
                tgt = 0
            # 专人专项的「至少上班天数」取该人自己的要求（覆盖班组目标）
            ov_days = int((worker_rules.get(p.name) or {}).get("work_days") or 0)
            if p.name in exempt_set:
                continue
            if ov_days > 0:
                # 该人有专项要求：以它为目标与硬性下限；不能再套用「豁免模式」的
                # min == max == 默认目标，否则会把专项要求夹死（无解）。
                req[p.name] = {"target": ov_days, "min": ov_days}
            elif tgt > 0:
                if has_exempt:
                    # 有豁免人员：非豁免恰好上满 tgt，剩余班次交给豁免人员平分
                    req[p.name] = {"target": tgt, "min": tgt, "max": tgt}
                else:
                    # 无豁免人员：非豁免最低 tgt（硬下限、不设上限），
                    # 多出来的班次由求解器均衡分配给这些人
                    req[p.name] = {"target": tgt, "min": tgt}
        return req

    config = dict(base_config)
    config["worker_shift_req"] = _build_req()
    result = build_schedule(config, time_limit_seconds=30, phase2_seconds=10)
    audit_log.info(
        "生成排班结束 team=%s status=%s reached=%s/%s user=%s",
        team.name, result.status,
        sum(1 for v in result.reached.values() if v), len(result.reached),
        request.user.username,
    )

    if not result.feasible:
        # 整体无解：创建记录保存诊断信息（无排班明细）。
        # 失败记录只保留「最新一条」用于看诊断——旧的失败记录必然没有明细，
        # 留着只会堆积成垃圾，这里一并清掉；成功的记录不受影响。
        Schedule.objects.filter(
            team=team, year=y, month=m
        ).exclude(status__in=_EFFECTIVE_STATUSES).delete()
        record = Schedule.objects.create(
            team=team, year=y, month=m, start_date=start, days=days,
            shifts=FIXED_SHIFTS, daily_total=team.daily_headcount,
            role_reqs=team.role_reqs or {}, min_shift_target=team.min_shift_target,
            exempt_names=team.exempt_names or [],
            rest_block={"min": rest_min, "max": rest_max},
            person_overrides=person_overrides,
            worker_snapshot=worker_snapshot,
            status=result.status, message=result.message, diagnostics=result.diagnostics,
        )
        return redirect("scheduler:schedule_result", pk=record.id)

    record = Schedule.objects.create(
        team=team, year=y, month=m, start_date=start, days=days,
        shifts=FIXED_SHIFTS, daily_total=team.daily_headcount,
        role_reqs=team.role_reqs or {},
        min_shift_target=team.min_shift_target,
        exempt_names=team.exempt_names or [],
        rest_block={"min": rest_min, "max": rest_max},
        person_overrides=person_overrides,
        worker_snapshot=worker_snapshot,
        status=result.status,
        message=result.message,
        diagnostics=result.diagnostics,
        single_rest=result.single_rest,
        rest_run_violations=result.rest_run_violations,
    )
    # 每个班组每个月只保留一份排班：删除该班组同月份更旧的排班（含其明细），
    # 保证班次展示、个人日历、排班记录三者数据一致
    old_schedules = Schedule.objects.filter(team=team, year=y, month=m).exclude(id=record.id)
    if old_schedules.exists():
        Assignment.objects.filter(schedule__in=old_schedules).delete()
        old_schedules.delete()
    # 明细统一写入 Assignment 表（唯一数据源）。
    # 当天实际岗位：按什么岗位上班就记什么岗位；岗位约束之外的补位人员记「普通」。
    name_map = {p.name: p for p in persons}
    assign_rows = []
    for d in range(days):
        for s in FIXED_SHIFTS:
            for nm in result.per_day[d][s]:
                p = name_map.get(nm)
                if p:
                    role = result.role_assignments.get((nm, d), "普通")
                    assign_rows.append(Assignment(
                        schedule=record, person=p, day=d, shift=s, role=role))
    Assignment.objects.bulk_create(assign_rows)
    return redirect("scheduler:schedule_result", pk=record.id)


# ---------------------------------------------------------------------------
# 人员编辑
# ---------------------------------------------------------------------------
@require_http_methods(["GET", "POST"])
@login_required
def person_edit(request, person_id):
    person = get_object_or_404(Person, id=person_id)
    ug = user_group(request)
    if ug and (person.team is None or person.team.group_id != ug.id):
        return redirect("scheduler:index")
    if not can_edit_team(request, person.team):
        return redirect("scheduler:person_detail", person_id=person.id)
    message = ""
    if request.method == "POST":
        selected = [rid for rid in (_post_int(v) for v in request.POST.getlist("roles")) if rid is not None]
        new_role_names = request.POST.get("new_roles", "")
        person.roles.set(Role.objects.filter(id__in=selected))
        for rn in re.split(r"[,，、\s]+", new_role_names.strip()):
            rn = (rn or "").strip()
            if rn and len(rn) <= 50:
                role, _ = Role.objects.get_or_create(name=rn)
                person.roles.add(role)
        team_id = _post_int(request.POST.get("team"))
        if team_id:
            tq = Team.objects.filter(id=team_id)
            if ug:
                tq = tq.filter(group=ug)
            new_team = tq.first()
            if new_team is not None:
                person.team = new_team
        else:
            person.team = None
        # 默认班次白名单：非法值保持原值，避免脏数据
        ds = (request.POST.get("default_shift") or "").strip()
        if ds in FIXED_SHIFTS:
            person.default_shift = ds
        person.is_active = request.POST.get("is_active") == "on"
        try:
            person.worked_so_far = max(0, min(999, int(request.POST.get("worked_so_far") or 0)))
            person.required_shifts = max(0, min(999, int(request.POST.get("required_shifts") or 0)))
        except ValueError:
            pass
        person.save()
        audit_log.info("编辑人员 person=%s team=%s shift=%s active=%s user=%s",
                       person.name, person.team_id, person.default_shift,
                       person.is_active, request.user.username)
        message = f"已保存「{person.name}」的信息。"
    # 岗位只显示「该人员所属队组」里出现过的岗位（并保留本人已有岗位），
    # 避免把全系统其它队组/导入误产生的无关岗位（乱码、人名等）列出来
    scope_group = person.team.group if person.team else ug
    if scope_group:
        all_roles = Role.objects.filter(
            Q(persons__team__group=scope_group) | Q(persons=person)
        ).distinct().order_by("name")
    else:
        all_roles = person.roles.all().order_by("name")

    return render(request, "scheduler/person_edit.html", {
        "person": person,
        "all_roles": all_roles,
        "teams": Team.objects.filter(group=ug) if ug else Team.objects.order_by("name"),
        "message": message,
    })


# ---------------------------------------------------------------------------
# 个人详情：日历 + 下井标签 + 改班（可按月份查看，默认 25 号起算的下一月）
# ---------------------------------------------------------------------------
# 「生成失败」的记录也会落库（为了在结果页展示诊断信息），但它没有任何明细。
# 因此凡是「取该班组该月生效排班」的地方，都必须优先取有明细的那一份；
# 否则一次失败的重新生成会把原本有效的排班从个人日历/班次展示上「藏掉」。
_EFFECTIVE_STATUSES = ("OPTIMAL", "FEASIBLE")


def _effective_schedule_qs(qs):
    """按「先成功的、再最新的」排序：成功记录优先于失败记录，同类按创建时间倒序。"""
    from django.db.models import Case, When, IntegerField
    return qs.annotate(
        _ok=Case(When(status__in=_EFFECTIVE_STATUSES, then=0),
                 default=1, output_field=IntegerField())
    ).order_by("_ok", "-created_at")


def _schedule_for(person: Person, year: int, month: int):
    """该人所属班组在指定年月的**生效**排班（优先有明细的成功记录）。"""
    if person.team_id is None:
        return None
    return _effective_schedule_qs(
        Schedule.objects.filter(team=person.team, year=year, month=month)
    ).first()


@login_required
def person_detail(request, person_id):
    person = get_object_or_404(Person, id=person_id)
    ug = user_group(request)
    if ug and (person.team is None or person.team.group_id != ug.id):
        return redirect("scheduler:index")
    message = ""
    error = ""

    # 月份参数；默认 = 25号起算的下一月周期
    y, m = _parse_year_month(request)

    schedule = _schedule_for(person, y, m)

    if request.method == "POST":
        if user_role(request) == "member":
            error = "队员账号为只读，不能修改排班。"
        else:
            try:
                day = int(request.POST.get("day", -1))
            except (TypeError, ValueError):
                day = -1
            new_shift = request.POST.get("shift", "")
            if not schedule:
                error = "该月还没有排班，无法改班。"
            elif not (0 <= day < schedule.days):
                error = "无效的改班请求。"
            else:
                ddate = schedule.start_date + timedelta(days=day) if schedule.start_date else None
                # 明细只存 Assignment 表：改班 = 增删改 Assignment，无需再同步结果 JSON
                if new_shift == "休息":
                    Assignment.objects.filter(schedule=schedule, person=person, day=day).delete()
                    message = f"已将{ddate:%m月%d日}改为休息。"
                elif new_shift in FIXED_SHIFTS:
                    assignment, _ = Assignment.objects.get_or_create(
                        schedule=schedule, person=person, day=day, defaults={"shift": new_shift})
                    assignment.shift = new_shift
                    assignment.save()
                    message = f"已将{ddate:%m月%d日}的班次改为「{new_shift}」。"
                else:
                    error = "无效的改班请求。"
                if not error:
                    audit_log.info("改班 user=%s person=%s schedule=%s day=%s -> %s",
                                   request.user.username, person.name, schedule.id,
                                   day, new_shift)

    start, end, days = period_range(y, m)
    assignments = {}
    if schedule:
        assignments = {a.day: a.shift for a in Assignment.objects.filter(schedule=schedule, person=person)}

    if request.GET.get("export") == "xlsx":
        return _export_person_xlsx(person, y, m, start, days, assignments)

    worked_auto = person_worked_auto(person, schedule)
    remaining = max(0, person.required_shifts - worked_auto)

    lead = start.weekday()
    cells = [None] * lead
    for d in range(days):
        cells.append({"index": d, "date": start + timedelta(days=d), "shift": assignments.get(d, "")})
    while len(cells) % 7 != 0:
        cells.append(None)
    weeks = [cells[i:i + 7] for i in range(0, len(cells), 7)]
    py, pm = month_shift(y, m, -1)
    ny, nm = month_shift(y, m, 1)

    return render(request, "scheduler/person_detail.html", {
        "person": person,
        "schedule": schedule,
        "year": y, "month": m,
        "prev_y": py, "prev_m": pm, "next_y": ny, "next_m": nm,
        "start": start, "end": end, "days": days,
        "today": date.today(),
        "weeks": weeks,
        "shifts": FIXED_SHIFTS,
        "worked_auto": worked_auto,
        "required": person.required_shifts,
        "remaining": remaining,
        "can_edit": can_edit_team(request, person.team),
        "message": message,
        "error": error,
    })


def _latest_schedules_for_month(year: int, month: int, group: Group = None):
    """某年月的班组排班（每个班组只取最新一份），可按队组过滤（保证班次展示与个人日历一致）。"""
    qs = Schedule.objects.filter(year=year, month=month)
    if group:
        qs = qs.filter(team__group=group)
    latest = {}
    for sch in _effective_schedule_qs(qs.select_related("team")):
        if sch.team_id is not None and sch.team_id not in latest:
            latest[sch.team_id] = sch
    return list(latest.values())


# ---------------------------------------------------------------------------
# 班次展示：日历显示每天每个班次人数（按班组分层）
# ---------------------------------------------------------------------------
@login_required
def shift_board(request):
    ug = user_group(request)
    groups = Group.objects.order_by("name")
    sel_group = request.GET.get("group", "")
    if ug:
        selected_group = ug
    else:
        selected_group = groups.filter(id=sel_group).first() if sel_group.isdigit() else None

    y, m = _parse_year_month(request)
    start, end, days = period_range(y, m)

    schedules = _latest_schedules_for_month(y, m, selected_group) if selected_group else []
    # board[day][shift][team_name] = [names]
    board = {d: {s: {} for s in FIXED_SHIFTS} for d in range(days)}
    for sch in schedules:
        per_day, _, _ = _assignment_matrix(sch)
        team_name = sch.team.name if sch.team else "未分组"
        for d in range(min(days, sch.days)):
            for s in FIXED_SHIFTS:
                names = per_day.get(d, {}).get(s, [])
                if names:
                    board[d][s][team_name] = names

    if request.GET.get("export") == "xlsx":
        return _export_board_xlsx(y, m, start, days, board)

    lead = start.weekday()
    cells = [None] * lead
    for d in range(days):
        day_total = 0
        shift_infos = []
        for s in FIXED_SHIFTS:
            teams_info = board[d][s]
            shift_infos.append({
                "name": s,
                "total": sum(len(n) for n in teams_info.values()),
                "teams": [(tname, len(names)) for tname, names in teams_info.items()],
            })
            day_total += shift_infos[-1]["total"]
        cells.append({"index": d, "date": start + timedelta(days=d),
                      "shifts": shift_infos, "total": day_total})
    while len(cells) % 7 != 0:
        cells.append(None)
    weeks = [cells[i:i + 7] for i in range(0, len(cells), 7)]
    py, pm = month_shift(y, m, -1)
    ny, nm = month_shift(y, m, 1)

    return render(request, "scheduler/shift_board.html", {
        "year": y, "month": m,
        "prev_y": py, "prev_m": pm, "next_y": ny, "next_m": nm,
        "start": start, "end": end, "days": days,
        "today": date.today(),
        "weeks": weeks,
        "shifts": FIXED_SHIFTS,
        "groups": groups,
        "selected_group": selected_group,
        "user_group": ug,
        "schedule_count": len(schedules),
        "can_add": user_role(request) in ("super", "team_admin"),
    })


@login_required
def shift_detail(request, year, month, day, shift):
    """某天某班次的详情：各班组上班人员（姓名/岗位/已上/应上），点击人名跳个人日历。"""
    ug = user_group(request)
    sel_group = request.GET.get("group", "")
    if ug:
        selected_group = ug
    else:
        selected_group = Group.objects.filter(id=sel_group).first() if sel_group.isdigit() else None
    if not (1 <= year <= 9999 and 1 <= month <= 12):
        return redirect("scheduler:shift_board")
    start, end, days = period_range(year, month)
    if not (0 <= day < days) or shift not in FIXED_SHIFTS:
        return redirect("scheduler:shift_board")
    ddate = start + timedelta(days=day)

    schedules = _latest_schedules_for_month(year, month, selected_group) if selected_group else []
    snap_by_team = {}
    names_today = []
    for sch in schedules:
        per_day, _, _ = _assignment_matrix(sch)
        names = per_day.get(day, {}).get(shift, [])
        if names:
            snap_by_team[sch.team.name if sch.team else "未分组"] = names
            names_today.extend(names)

    # 当天实际岗位：按什么岗位上班就显示什么岗位（旧数据无岗位记录时回退为空，前端显示人员岗位）
    role_of_day = dict(Assignment.objects.filter(
        schedule__in=schedules, day=day, shift=shift, person__name__in=names_today
    ).values_list("person__name", "role"))

    persons = {p.name: p for p in Person.objects.prefetch_related("roles").filter(name__in=names_today)}
    schedule_map = {}
    for sch in schedules:
        if sch.team_id not in schedule_map:
            schedule_map[sch.team_id] = sch

    # 批量计算已上班数（一次查询，避免逐人 N+1）
    pairs = []
    for team_name in sorted(snap_by_team):
        for nm in snap_by_team[team_name]:
            p = persons.get(nm)
            if p and p.team_id in schedule_map:
                pairs.append((p, schedule_map[p.team_id]))
    worked_map = _worked_auto_batch(pairs)

    groups = []
    for team_name in sorted(snap_by_team):
        entries = []
        for nm in snap_by_team[team_name]:
            p = persons.get(nm)
            entries.append({
                "person": p,
                "person_name": nm,
                "role": role_of_day.get(nm, ""),
                "worked": worked_map.get(p.id, p.worked_so_far) if p else 0,
                "required": p.required_shifts if p else 0,
            })
        groups.append({"team": team_name, "entries": entries})

    return render(request, "scheduler/shift_detail.html", {
        "year": year, "month": month, "day": day, "shift": shift,
        "date": ddate,
        "groups": groups,
        "selected_group": selected_group,
        "is_member": user_role(request) == "member",
    })


@login_required
def shift_add(request, year, month, day, shift):
    """某天某班次的「加人」：列出当天休息（未上班）的人员，按班组分组，点选加入该班次。"""
    if user_role(request) == "member":
        return redirect("scheduler:shift_board")
    ug = user_group(request)
    sel_group = request.GET.get("group", "")
    if ug:
        selected_group = ug
    else:
        selected_group = Group.objects.filter(id=sel_group).first() if sel_group.isdigit() else None

    if not (1 <= year <= 9999 and 1 <= month <= 12):
        return redirect("scheduler:shift_board")
    start, end, days = period_range(year, month)
    if not (0 <= day < days) or shift not in FIXED_SHIFTS:
        return redirect("scheduler:shift_board")
    ddate = start + timedelta(days=day)
    gid = selected_group.id if selected_group else ""

    message = ""
    error = ""
    if request.method == "POST":
        pid = _post_int(request.POST.get("person_id"))
        person = Person.objects.select_related("team").filter(id=pid).first() if pid else None
        if not person:
            error = "人员不存在或已被删除。"
        elif person.team is None:
            error = "该人员未分组，无法加人。"
        elif ug and person.team.group_id != ug.id:
            error = "只能给本队组的人员加人。"
        else:
            schedule = _schedule_for(person, year, month)
            if not schedule:
                error = f"「{person.team.name}」本月还没有排班，无法加人。请先生成排班。"
            else:
                assignment, _ = Assignment.objects.get_or_create(
                    schedule=schedule, person=person, day=day,
                    defaults={"shift": shift, "role": "普通"})
                assignment.shift = shift
                assignment.save()
                audit_log.info("加人 user=%s person=%s schedule=%s day=%s shift=%s",
                               request.user.username, person.name, schedule.id, day, shift)
                from urllib.parse import quote
                return redirect(f"{request.path}?group={gid}&added={quote(person.name)}")

    added_name = request.GET.get("added", "")
    if added_name:
        message = f"已把「{added_name}」加为 {ddate:%m月%d日} 的{shift}。"

    groups = []
    schedules = _latest_schedules_for_month(year, month, selected_group) if selected_group else []
    for sch in schedules:
        team = sch.team
        if not team:
            continue
        persons = list(Person.objects.filter(team=team, is_active=True).order_by("name"))
        if not persons:
            continue
        working_ids = set(Assignment.objects.filter(schedule=sch, day=day).values_list("person_id", flat=True))
        resting = [p for p in persons if p.id not in working_ids]
        if resting:
            groups.append({"team": team.name, "resting": resting})

    return render(request, "scheduler/shift_add.html", {
        "year": year, "month": month, "day": day, "shift": shift,
        "date": ddate,
        "groups": groups,
        "selected_group": selected_group,
        "message": message,
        "error": error,
    })


# ---------------------------------------------------------------------------
# 排班结果 / 记录
# ---------------------------------------------------------------------------
@login_required
def schedule_result(request, pk):
    record = get_object_or_404(Schedule, id=pk)
    ug = user_group(request)
    if ug and (record.team is None or record.team.group_id != ug.id):
        return redirect("scheduler:index")

    shifts = record.shifts or FIXED_SHIFTS
    # 明细统一从 Assignment 表聚合（唯一数据源）
    per_day, wc, _ = _assignment_matrix(record)
    snap = {s["name"]: s for s in (record.worker_snapshot or [])}
    # worker_counts 包含所有人（无班的人计 0）
    worker_counts = {nm: 0 for nm in snap}
    worker_counts.update(wc)

    # 实时重算达标 / 未达标（规则3：全局「至少应上班数」优先，否则每人「应上班数」）
    exempt = set(record.exempt_names or [])
    target_global = record.min_shift_target or 0
    reached = {}
    shortfall = []
    for nm, s in snap.items():
        if nm in exempt:
            continue
        if target_global > 0:
            tgt = target_global - s["worked"]
        elif s["required"] > 0:
            tgt = s["required"] - s["worked"]
        else:
            tgt = 0
        if tgt <= 0:
            continue
        cnt = worker_counts.get(nm, 0)
        reached[nm] = cnt >= tgt
        if cnt < tgt:
            shortfall.append({"name": nm, "target": tgt, "count": cnt})
    shortfall.sort(key=lambda x: x["target"] - x["count"], reverse=True)

    rows = []
    for d in range(1, record.days + 1):
        ddate = record.start_date + timedelta(days=d - 1) if record.start_date else None
        cells = per_day.get(d - 1, {})
        rows.append({"day": d, "date": ddate, "cells": [(s, cells.get(s, [])) for s in shifts]})

    def _status(nm):
        if nm in reached:
            return "ok" if reached[nm] else "no"
        return "exempt"

    count_rows = sorted(
        [(nm, cnt, _status(nm),
          snap.get(nm, {}).get("worked", 0),
          snap.get(nm, {}).get("required", 0))
         for nm, cnt in worker_counts.items()],
        key=lambda kv: (0 if kv[2] == "ok" else 1, -kv[1]),
    )

    capacity = None
    if record.team:
        people_count = len(worker_counts)
        exempt_count = len([n for n in (record.exempt_names or []) if n in worker_counts])
        capacity = capacity_quick(
            people_count, record.daily_total or 0, record.days,
            (record.rest_block or {}).get("max", 4),
            target_global if target_global > 0 else 18,
        )
        reached_count = sum(1 for v in reached.values() if v)
        capacity["max_fillable"] = reached_count
        capacity["needed_exempt"] = max(0, (people_count - exempt_count) - reached_count)
        capacity["people"] = people_count
        capacity["daily"] = record.daily_total or 0
        capacity["target"] = target_global
    return render(request, "scheduler/schedule_result.html", {
        "record": record,
        "shifts": shifts,
        "rows": rows,
        "count_rows": count_rows,
        "shortfall": shortfall,
        "capacity": capacity,
        "stats": {
            "total": sum(worker_counts.values()),
            "avg": round(sum(worker_counts.values()) / len(worker_counts), 1)
            if worker_counts else 0,
            "reached_count": sum(1 for v in reached.values() if v),
            "reached_total": len(reached),
            "single_rest": record.single_rest,
            "rest_violations": record.rest_run_violations,
            "target": target_global or None,
        },
    })


@login_required
def schedule_list(request):
    # 队员（只读）不可访问排班记录
    if user_role(request) == "member":
        return redirect("scheduler:index")
    ug = user_group(request)
    records = Schedule.objects.select_related("team").order_by("-created_at")
    if ug:
        records = records.filter(team__group=ug)
    return render(request, "scheduler/schedule_list.html", {"records": records})
